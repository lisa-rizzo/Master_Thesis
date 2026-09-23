"""
Truncated epoch analysis: GLARE and GLAREx detection performance using only first N epochs.

Reuses the existing 8-epoch grad_norms file without retraining.
Answers: how many checkpoint epochs do GLARE/GLAREx actually need?

For each N in 1..8:
  - Filter grad_norms to epochs 0..(N-1)
  - Recompute GLARE and GLAREx scores via calculate_glare()
  - GLARE : integer threshold table (recall, precision, F1 at each integer threshold)
  - GLAREx: AUROC and recall curves only (float score, not suited for integer threshold tables)

The distinction:
  - GLARE  score is an integer in [0, N-1], naturally swept at thresholds 0,1,...,N-1
  - GLAREx score is a float refinement; within the same GLARE bucket it re-ranks by gradient
    magnitude, but the 0.01*(1/avg) term is unbounded (can exceed 1), so integer thresholds
    do not correspond to the same sample sets as GLARE thresholds -> use AUROC and recall
    curves instead of a threshold table for GLAREx.
"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent))

from datetime import datetime
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from sklearn.metrics import roc_auc_score
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

from glare import calculate_glare

# ── Paths ─────────────────────────────────────────────────────────────────────
GRAD_NORMS_PKL = Path("/home/student/lisa_ma/results/glare/19-08-26_21:11_ep-8_single-gpu/results/grad_norms_20260819_211102.pkl")
GT_XLSX        = Path("/home/student/lisa_ma/evaluation/gt_mislabels_complete.xlsx")
OUTPUT_DIR     = Path("/home/student/lisa_ma/results/early_epoch_comparison")
FIGURES_DIR    = Path("/home/student/lisa_ma/results/early_epoch_comparison")


# ── Metric helpers ─────────────────────────────────────────────────────────────
def threshold_metrics_glare(y_true, scores_int, threshold):
    """Threshold table row for integer GLARE scores."""
    predicted = (scores_int <= threshold).astype(int)
    tp = int(((predicted == 1) & (y_true == 1)).sum())
    fp = int(((predicted == 1) & (y_true == 0)).sum())
    fn = int(((predicted == 0) & (y_true == 1)).sum())
    tn = int(((predicted == 0) & (y_true == 0)).sum())
    recall    = tp / (tp + fn) if (tp + fn) > 0 else 0.0
    precision = tp / (tp + fp) if (tp + fp) > 0 else 0.0
    f1        = 2 * precision * recall / (precision + recall) if (precision + recall) > 0 else 0.0
    fpr       = fp / (fp + tn) if (fp + tn) > 0 else 0.0
    return dict(
        threshold=threshold,
        samples_flagged=tp + fp,
        frac_checked_pct=round((tp + fp) / len(y_true) * 100, 2),
        recall_pct=round(recall * 100, 1),
        precision_pct=round(precision * 100, 1),
        f1_pct=round(f1 * 100, 1),
        fpr_pct=round(fpr * 100, 2),
        TP=tp, FP=fp, FN=fn, TN=tn,
    )


def compute_glare_table(scores_df, positives, n_epochs):
    """Integer threshold table for GLARE scores."""
    y_true = scores_df["id"].isin(positives).astype(int).values
    scores = scores_df["glare score"].values.astype(int)
    rows = [threshold_metrics_glare(y_true, scores, k) for k in range(n_epochs)]
    return pd.DataFrame(rows)


def compute_auroc(scores_df, positives, score_col):
    """AUROC for any score column (lower = more suspicious)."""
    y_true   = scores_df["id"].isin(positives).astype(int).values
    scores   = scores_df[score_col].values
    inverted = scores.max() - scores
    try:
        return roc_auc_score(y_true, inverted) if len(set(y_true)) > 1 else float("nan")
    except Exception:
        return float("nan")


def recall_curve_sorted(scores_df, positives, score_col):
    """Recall vs fraction-checked by sorting on score (ascending = most suspicious first)."""
    df = scores_df.copy()
    df["is_mislabel"] = df["id"].isin(positives).astype(int)
    df_sorted = df.sort_values(score_col, ascending=True).reset_index(drop=True)

    total     = len(df_sorted)
    total_mis = df_sorted["is_mislabel"].sum()

    fracs, recalls = [0.0], [0.0]
    cum_tp = 0
    for i, row in df_sorted.iterrows():
        cum_tp += row["is_mislabel"]
        fracs.append((i + 1) / total)
        recalls.append(cum_tp / total_mis if total_mis > 0 else 0.0)
    return fracs, recalls


# ── Main ───────────────────────────────────────────────────────────────────────
def main():
    print("Loading grad norms...")
    gn = pd.read_pickle(GRAD_NORMS_PKL)

    print("Loading ground truth...")
    df_gt     = pd.read_excel(GT_XLSX, dtype=str)
    positives = set(df_gt["sample_id"].astype(str))

    total_epochs = int(gn["epoch"].max()) + 1
    n_values     = list(range(1, total_epochs + 1))

    print(f"Total epochs available : {total_epochs}")
    print(f"GT positives (train)   : {len(positives)}")
    print(f"Total samples          : {gn['id'].nunique()}\n")

    # ── Loop over N ───────────────────────────────────────────────────────────
    summary_rows      = []
    glare_tables      = {}   # n -> df_metrics (GLARE integer threshold table)
    recall_curves_all = {}   # (n, method) -> (fracs, recalls)

    for n in n_values:
        subset    = gn[gn["epoch"] < n].copy()
        scores_df = calculate_glare(subset)

        # GLARE: integer threshold table + AUROC
        df_glare_tbl = compute_glare_table(scores_df, positives, n)
        glare_tables[n] = df_glare_tbl
        auroc_glare  = compute_auroc(scores_df, positives, "glare score")
        auroc_glarex = compute_auroc(scores_df, positives, "glarex score")

        # Recall curves
        recall_curves_all[(n, "GLARE")]  = recall_curve_sorted(scores_df, positives, "glare score")
        recall_curves_all[(n, "GLAREx")] = recall_curve_sorted(scores_df, positives, "glarex score")

        # Summary row
        row = {"N_epochs": n, "AUROC_GLARE": round(auroc_glare, 4), "AUROC_GLAREx": round(auroc_glarex, 4)}
        for k in range(min(n, 8)):
            thr_row = df_glare_tbl[df_glare_tbl["threshold"] == k]
            if len(thr_row):
                r = thr_row.iloc[0]
                row[f"Recall_thr{k}"]    = r["recall_pct"]
                row[f"Precision_thr{k}"] = r["precision_pct"]
                row[f"F1_thr{k}"]        = r["f1_pct"]
                row[f"Flagged_thr{k}"]   = r["samples_flagged"]
        summary_rows.append(row)
        print(f"N={n:2d}  GLARE AUROC={auroc_glare:.4f}  |  GLAREx AUROC={auroc_glarex:.4f}")

    df_summary = pd.DataFrame(summary_rows)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    FIGURES_DIR.mkdir(exist_ok=True)

    # ── Plot 1: AUROC vs N (GLARE and GLAREx side-by-side) ───────────────────
    fig, axes = plt.subplots(1, 2, figsize=(13, 5))

    for ax, method, col in zip(axes, ["GLARE", "GLAREx"], ["AUROC_GLARE", "AUROC_GLAREx"]):
        yvals = df_summary[col].values
        ax.plot(n_values, yvals, "o-", linewidth=2, markersize=6, color="#2471A3")
        ax.axhline(yvals[-1], color="gray", linewidth=0.8, linestyle="--", alpha=0.6,
                   label=f"Full {total_epochs} epochs ({yvals[-1]:.3f})")
        ax.set_xlabel("Number of Epochs Used", fontsize=11)
        ax.set_ylabel("AUROC", fontsize=11)
        ax.set_xticks(n_values)
        ax.set_ylim(max(0.5, yvals.min() - 0.05), min(1.0, yvals.max() + 0.02))
        ax.set_title(f"{method}: AUROC vs. Epochs Used", fontsize=12)
        ax.legend(fontsize=9)
        ax.grid(True, linestyle="--", alpha=0.3)
        for sp in ax.spines.values():
            sp.set_linewidth(0.6)
            sp.set_color("#AAAAAA")

    plt.tight_layout()
    fig.savefig(str(FIGURES_DIR / f"truncated_epoch_auroc_{ts}.pdf"), bbox_inches="tight")
    fig.savefig(str(FIGURES_DIR / f"truncated_epoch_auroc_{ts}.png"), dpi=180, bbox_inches="tight")
    print(f"\nFigure: truncated_epoch_auroc_{ts}.pdf")

    # ── Plot 2: Recall curves (GLARE vs GLAREx, zoom 0–15% checked) ──────────
    fig2, axes2 = plt.subplots(1, 2, figsize=(14, 5))
    cmap = plt.cm.plasma

    for ax, method in zip(axes2, ["GLARE", "GLAREx"]):
        for i, n in enumerate(n_values):
            color = cmap(i / max(len(n_values) - 1, 1))
            fracs, recalls = recall_curves_all[(n, method)]
            auroc_col = "AUROC_GLARE" if method == "GLARE" else "AUROC_GLAREx"
            auroc_val = df_summary.loc[df_summary["N_epochs"] == n, auroc_col].values[0]
            lw = 2.2 if n == total_epochs else 1.2
            ls = "-"  if n == total_epochs else "--"
            ax.plot(fracs, recalls, color=color, linewidth=lw, linestyle=ls,
                    label=f"N={n} (AUROC {auroc_val:.3f})")
        ax.set_xlabel("Fraction of Training Data Checked", fontsize=11)
        ax.set_ylabel("Recall", fontsize=11)
        ax.set_xlim(0, 0.15)
        ax.set_ylim(0, 1.05)
        ax.legend(fontsize=7.5, loc="lower right", framealpha=0.95,
                  edgecolor="#CCCCCC", fancybox=False)
        ax.set_title(f"{method}: Recall vs. Fraction Checked", fontsize=12)
        for sp in ax.spines.values():
            sp.set_linewidth(0.6)
            sp.set_color("#AAAAAA")

    plt.tight_layout()
    fig2.savefig(str(FIGURES_DIR / f"truncated_epoch_recall_{ts}.pdf"), bbox_inches="tight")
    fig2.savefig(str(FIGURES_DIR / f"truncated_epoch_recall_{ts}.png"), dpi=180, bbox_inches="tight")
    print(f"Figure: truncated_epoch_recall_{ts}.pdf")

    # ── Excel ─────────────────────────────────────────────────────────────────
    out_path = OUTPUT_DIR / f"truncated_epoch_analysis_{ts}.xlsx"

    HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True)
    GREEN       = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")

    def style_sheet(ws, n_data_rows):
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")
        header = [c.value for c in ws[1]]
        if "F1 (%)" in header:
            f1_col = header.index("F1 (%)") + 1
            f1_vals = [ws.cell(row=r, column=f1_col).value for r in range(2, n_data_rows + 2)]
            valid = [v for v in f1_vals if v is not None]
            if valid:
                best = f1_vals.index(max(valid)) + 2
                for cell in ws[best]:
                    cell.fill = GREEN
        for col in ws.columns:
            ws.column_dimensions[get_column_letter(col[0].column)].width = 16
        ws.freeze_panes = "A2"

    output_cols  = ["threshold", "samples_flagged", "frac_checked_pct",
                    "recall_pct", "precision_pct", "f1_pct",
                    "TP", "FP", "FN", "TN"]
    display_cols = ["Threshold (<=)", "Flagged", "Fraction checked (%)",
                    "Recall (%)", "Precision (%)", "F1 (%)",
                    "TP", "FP", "FN", "TN"]

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df_summary.to_excel(writer, sheet_name="Summary_AUROC", index=False)
        for n in n_values:
            df_out = glare_tables[n][output_cols].copy()
            df_out.columns = display_cols
            df_out.to_excel(writer, sheet_name=f"GLARE_N{n}", index=False)

    wb = load_workbook(out_path)
    for n in n_values:
        style_sheet(wb[f"GLARE_N{n}"], n)
    wb.save(out_path)

    print(f"\nExcel: {out_path}")
    print("\nAUROC summary:")
    print(df_summary[["N_epochs", "AUROC_GLARE", "AUROC_GLAREx"]].to_string(index=False))


if __name__ == "__main__":
    main()
