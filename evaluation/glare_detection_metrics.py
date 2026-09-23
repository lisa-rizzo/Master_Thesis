"""
Evaluate GLARE detection performance on region-crossing mislabels.

Positives: region-crossing mislabels from gt_mislabels_complete.xlsx, filtered to
           training-split PIDs only (GLARE only scores training samples).
Negatives: everything else in the full mislabels_list CSV.

Outputs:
  - Per-threshold metrics table (recall, precision, F1, FPR, fraction checked)
  - AUROC and Average Precision
  - Excel with the metrics table + a ranked list of the lowest-scoring samples
"""
from pathlib import Path
import logging
import pandas as pd
import numpy as np
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

logging.basicConfig(level=logging.INFO, format="%(message)s")
log = logging.getLogger("glare_metrics")

MISLABELS_CSV  = Path("/home/student/lisa_ma/results/glare/22-08-26_11:28_ep-15_dual-gpu/results/mislabels_list_20260822_112835.csv")
GT_EXCEL       = Path("/home/student/lisa_ma/evaluation/gt_mislabels_complete.xlsx")
DATA_FILTER    = Path("/home/student/lisa_ma/datasets/VerSe/data_filter_joined.xlsx")
OUTPUT_DIR     = Path("/home/student/lisa_ma/evaluation")

MAX_SCORE      = 15  # number of training epochs


def compute_auroc_ap(y_true, scores):
    """AUROC and Average Precision from GLARE scores (lower = more suspicious)."""
    from sklearn.metrics import roc_auc_score, average_precision_score
    # Invert scores so that higher = more suspicious for sklearn
    inverted = MAX_SCORE - scores
    auroc = roc_auc_score(y_true, inverted)
    ap    = average_precision_score(y_true, inverted)
    return auroc, ap


def threshold_metrics(y_true, scores, threshold):
    predicted = (scores <= threshold).astype(int)
    tp = int(((predicted == 1) & (y_true == 1)).sum())
    fp = int(((predicted == 1) & (y_true == 0)).sum())
    fn = int(((predicted == 0) & (y_true == 1)).sum())
    tn = int(((predicted == 0) & (y_true == 0)).sum())

    recall    = tp / (tp + fn) if (tp + fn) > 0 else 0.0
    precision = tp / (tp + fp) if (tp + fp) > 0 else 0.0
    f1        = 2 * precision * recall / (precision + recall) if (precision + recall) > 0 else 0.0
    fpr       = fp / (fp + tn) if (fp + tn) > 0 else 0.0
    frac_checked = (tp + fp) / len(y_true)

    return dict(
        threshold=threshold,
        TP=tp, FP=fp, FN=fn, TN=tn,
        recall=round(recall, 4),
        precision=round(precision, 4),
        F1=round(f1, 4),
        FPR=round(fpr, 4),
        fraction_checked=round(frac_checked, 4),
        samples_flagged=tp + fp,
    )


def main():
    log.info("GT mislabels:  %s", GT_EXCEL.name)
    log.info("Mislabels CSV: %s", MISLABELS_CSV.name)

    # Load ground-truth mislabels and filter to training split only
    df_gt = pd.read_excel(GT_EXCEL, dtype=str)
    df_filter = pd.read_excel(DATA_FILTER, dtype=str)
    train_pids = set(df_filter.loc[df_filter["split"] == "train", "pid"].str.strip())
    df_gt_train = df_gt[df_gt["pid"].str.strip().isin(train_pids)]
    positives = set(df_gt_train["sample_id"].astype(str))
    log.info("Region-crossing positives (in training): %d", len(positives))

    # Load full dataset
    df_full = pd.read_csv(MISLABELS_CSV, dtype=str, keep_default_na=False)
    df_full["glare_score"] = pd.to_numeric(df_full["glare score"], errors="coerce")
    df_full = df_full.dropna(subset=["glare_score"])

    total = len(df_full)
    log.info("Total samples with GLARE score: %d", total)

    # Binary labels
    df_full["is_mislabel"] = df_full["id"].isin(positives).astype(int)
    n_pos = df_full["is_mislabel"].sum()
    n_neg = total - n_pos
    log.info("Positives: %d  |  Negatives: %d", n_pos, n_neg)

    y_true  = df_full["is_mislabel"].values
    scores  = df_full["glare_score"].values

    # ── Per-exact-score counts ───────────────────────────────────────────────
    exact_counts = {}
    for s in range(0, MAX_SCORE + 1):
        mask = scores == s
        exact_counts[s] = {
            "n_samples_exact_score":   int(mask.sum()),
            "n_mislabels_exact_score": int((mask & (y_true == 1)).sum()),
        }

    # ── Threshold sweep ──────────────────────────────────────────────────────
    thresholds = list(range(0, MAX_SCORE + 1))
    rows = [threshold_metrics(y_true, scores, t) for t in thresholds]
    df_metrics = pd.DataFrame(rows)

    # Attach exact-score columns
    df_metrics["n_samples_exact_score"]   = df_metrics["threshold"].map(
        lambda t: exact_counts[t]["n_samples_exact_score"])
    df_metrics["n_mislabels_exact_score"] = df_metrics["threshold"].map(
        lambda t: exact_counts[t]["n_mislabels_exact_score"])

    log.info("\n%s", "=" * 70)
    log.info("Threshold | TP  FP    FN  TN     | Recall  Prec   F1     FPR    | Frac.checked")
    log.info("-" * 70)
    for _, r in df_metrics.iterrows():
        log.info(
            "  <= %-4d  | %-3d %-5d %-3d %-5d | %-7.1f%% %-6.1f%% %-6.1f%% %-6.1f%% | %.2f%%",
            r.threshold, r.TP, r.FP, r.FN, r.TN,
            r.recall*100, r.precision*100, r.F1*100, r.FPR*100,
            r.fraction_checked*100,
        )

    # ── AUROC / AP ───────────────────────────────────────────────────────────
    try:
        auroc, ap = compute_auroc_ap(y_true, scores)
        log.info("\nAUROC: %.4f  |  Average Precision: %.4f", auroc, ap)
    except Exception as e:
        auroc, ap = float("nan"), float("nan")
        log.warning("Could not compute AUROC/AP: %s", e)

    # ── Ranked list of lowest-scoring samples ────────────────────────────────
    df_ranked = df_full.sort_values("glare_score")[
        ["id", "label", "glare score", "alternative class (glare)", "is_mislabel"]
    ].copy()
    df_ranked.columns = ["id", "training_class_code", "glare_score",
                         "alternative_class_code", "is_true_positive"]
    df_ranked["is_true_positive"] = df_ranked["is_true_positive"].astype(bool)

    # ── Write Excel ──────────────────────────────────────────────────────────
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = OUTPUT_DIR / f"glare_detection_metrics_{ts}.xlsx"

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # Metrics sheet
        df_metrics_out = df_metrics.copy()
        df_metrics_out["recall_%"]   = (df_metrics_out["recall"]   * 100).round(1)
        df_metrics_out["precision_%"] = (df_metrics_out["precision"] * 100).round(1)
        df_metrics_out["F1_%"]       = (df_metrics_out["F1"]       * 100).round(1)
        df_metrics_out["FPR_%"]      = (df_metrics_out["FPR"]      * 100).round(1)
        df_metrics_out["frac_checked_%"] = (df_metrics_out["fraction_checked"] * 100).round(2)
        cols = ["threshold", "n_samples_exact_score", "n_mislabels_exact_score",
                "TP", "FP", "FN", "TN",
                "recall_%", "precision_%", "F1_%", "FPR_%", "frac_checked_%", "samples_flagged"]
        df_metrics_out[cols].to_excel(writer, sheet_name="Threshold_Metrics", index=False)

        # Summary sheet
        summary = pd.DataFrame([
            ("Total samples", total),
            ("Region-crossing positives (in training)", n_pos),
            ("Negatives", n_neg),
            ("Base rate (%)", round(100 * n_pos / total, 4)),
            ("", ""),
            ("AUROC", round(auroc, 4)),
            ("Average Precision", round(ap, 4)),
            ("Max score (epochs)", MAX_SCORE),
        ], columns=["Metric", "Value"])
        summary.to_excel(writer, sheet_name="Summary", index=False)

        # Ranked list (top 500 most suspicious)
        df_ranked.head(500).to_excel(writer, sheet_name="Top500_Suspicious", index=False)

    # ── Formatting ────────────────────────────────────────────────────────────
    wb = load_workbook(out_path)
    ws = wb["Threshold_Metrics"]

    HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True)
    GREEN       = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
    YELLOW      = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    # Highlight best F1 row
    f1_col_idx = cols.index("F1_%") + 1
    f1_vals = [ws.cell(row=r, column=f1_col_idx).value for r in range(2, len(rows) + 2)]
    best_row = f1_vals.index(max(f1_vals)) + 2
    for cell in ws[best_row]:
        cell.fill = GREEN

    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = 18
    ws.freeze_panes = "A2"

    # Color true positives green in ranked sheet
    ws_rank = wb["Top500_Suspicious"]
    for cell in ws_rank[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    tp_col = 5  # is_true_positive
    for row in ws_rank.iter_rows(min_row=2):
        if row[tp_col - 1].value:
            for cell in row:
                cell.fill = GREEN

    wb.save(out_path)
    log.info("Output: %s", out_path)


if __name__ == "__main__":
    main()
